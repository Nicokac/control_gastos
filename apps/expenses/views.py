import hashlib
import io
import json
import logging
import tempfile
from decimal import Decimal, InvalidOperation

from django.contrib import messages
from django.contrib.auth.mixins import LoginRequiredMixin
from django.db import transaction
from django.db.models import Q, Sum
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.urls import reverse_lazy
from django.utils import timezone
from django.views import View

from apps.categories.models import Category
from apps.core.constants import CategoryType, PaymentMethod
from apps.core.utils import get_month_date_range_exclusive
from apps.core.views import (
    UserOwnedCreateView,
    UserOwnedDeleteView,
    UserOwnedDetailView,
    UserOwnedListView,
    UserOwnedUpdateView,
)

from .forms import ExpenseFilterForm, ExpenseForm
from .models import Expense

logger = logging.getLogger(__name__)


class ExpenseListView(UserOwnedListView):
    model = Expense
    template_name = "expenses/expense_list.html"
    context_object_name = "expenses"

    def get_queryset(self):
        qs = super().get_queryset().select_related("category", "category__parent", "saving")

        has_filters = any(
            key in self.request.GET
            for key in [
                "q",
                "month",
                "year",
                "category",
                "subcategory",
                "date_from",
                "date_to",
                "payment_method",
            ]
        )

        if has_filters:
            month = self.request.GET.get("month")
            year = self.request.GET.get("year")
        else:
            today = timezone.localdate()
            month = str(today.month)
            year = str(today.year)

        q = self.request.GET.get("q", "").strip()
        category = self.request.GET.get("category")
        subcategory = self.request.GET.get("subcategory")
        payment_method = self.request.GET.get("payment_method")

        # Filtro por fecha exacta (date_from / date_to) — tiene prioridad sobre mes/año
        date_from = self.request.GET.get("date_from")
        date_to = self.request.GET.get("date_to")
        if date_from or date_to:
            try:
                from datetime import date as date_cls

                if date_from:
                    qs = qs.filter(date__gte=date_cls.fromisoformat(date_from))
                if date_to:
                    qs = qs.filter(date__lte=date_cls.fromisoformat(date_to))
            except ValueError:
                pass
        elif month and year:
            # ✅ Mes/año -> rango [start, end)
            try:
                month_int = int(month)
                year_int = int(year)
                if 1 <= month_int <= 12 and 1900 <= year_int <= 2100:
                    start, end = get_month_date_range_exclusive(month_int, year_int)
                    qs = qs.filter(date__gte=start, date__lt=end)
            except ValueError:
                pass
        else:
            # Si viene solo year, mantenemos comportamiento actual
            if year:
                try:
                    year_int = int(year)
                    if 1900 <= year_int <= 2100:
                        qs = qs.filter(date__year=year_int)
                except ValueError:
                    pass

        if q:
            qs = qs.filter(
                Q(description__icontains=q)
                | Q(category__name__icontains=q)
                | Q(category__parent__name__icontains=q)
            )
        if subcategory:
            # subcategoría específica tiene prioridad sobre el grupo
            qs = qs.filter(category_id=subcategory)
        elif category:
            # category contiene el pk de un grupo (parent); filtramos sus subcategorías
            qs = qs.filter(category__parent_id=category)
        if payment_method:
            qs = qs.filter(payment_method=payment_method)

        order_by = self.request.GET.get("order_by", "date")
        direction = self.request.GET.get("dir", "desc")
        allowed_fields = {
            "date": "date",
            "category": "category__name",
            "description": "description",
            "amount": "amount_ars",
        }
        field = allowed_fields.get(order_by, "date")
        prefix = "-" if direction == "desc" else ""
        qs = qs.order_by(f"{prefix}{field}", "-created_at")
        return qs

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)

        has_filters = any(
            key in self.request.GET
            for key in [
                "q",
                "month",
                "year",
                "category",
                "subcategory",
                "date_from",
                "date_to",
                "payment_method",
            ]
        )

        if has_filters:
            form_data = self.request.GET
        else:
            today = timezone.localdate()
            form_data = {"month": today.month, "year": today.year}

        context["filter_form"] = ExpenseFilterForm(form_data, user=self.request.user)
        context["has_active_filters"] = any(
            self.request.GET.get(key)
            for key in [
                "q",
                "category",
                "subcategory",
                "payment_method",
            ]
        )
        context["order_by"] = self.request.GET.get("order_by", "date")
        context["order_dir"] = self.request.GET.get("dir", "desc")

        qs = self.object_list

        total = qs.aggregate(total=Sum("amount_ars"))["total"] or 0
        context["total"] = total
        total_nonzero = total or 1

        payment_method_labels = dict(PaymentMethod.choices)
        method_classified = qs.exclude(payment_method="").aggregate(s=Sum("amount_ars"))["s"] or 0
        method_unclassified = (total - method_classified) if total else 0
        payment_method_summary = [
            {
                "label": payment_method_labels.get(row["payment_method"], row["payment_method"]),
                "subtotal": row["subtotal"],
            }
            for row in qs.exclude(payment_method="")
            .values("payment_method")
            .annotate(subtotal=Sum("amount_ars"))
            .order_by("payment_method")
        ]
        if method_unclassified > 0:
            payment_method_summary.append(
                {"label": "Sin clasificar", "subtotal": method_unclassified}
            )
        _method_colors = {
            "Efectivo": "#28a745",
            "Débito": "#0dcaf0",
            "Crédito": "#dc3545",
            "Transferencia": "#6610f2",
            "Sin clasificar": "#adb5bd",
        }
        for item in payment_method_summary:
            item["pct"] = round(item["subtotal"] / total_nonzero * 100, 1)
            item["color"] = _method_colors.get(item["label"], "#6c757d")
        payment_method_summary.sort(key=lambda x: x["subtotal"], reverse=True)
        context["payment_method_summary"] = payment_method_summary
        context["method_donut_labels"] = [i["label"] for i in payment_method_summary]
        context["method_donut_data"] = [float(i["subtotal"]) for i in payment_method_summary]
        context["method_donut_colors"] = [i["color"] for i in payment_method_summary]

        # Agrupar por grupo (parent), no por subcategoría — fuente única para donut y leyenda
        group_totals = {}
        for row in qs.values(
            "category__parent_id",
            "category__parent__name",
            "category__parent__color",
            "category_id",
            "category__name",
            "category__color",
        ).annotate(subtotal=Sum("amount_ars")):
            if row["category__parent_id"]:
                gid = row["category__parent_id"]
                gname = row["category__parent__name"]
                gcolor = row["category__parent__color"] or "#6c757d"
            else:
                gid = row["category_id"]
                gname = row["category__name"]
                gcolor = row["category__color"] or "#6c757d"

            if gid not in group_totals:
                group_totals[gid] = {"name": gname, "color": gcolor, "pk": gid, "subtotal": 0}
            group_totals[gid]["subtotal"] += row["subtotal"]

        group_summary = sorted(group_totals.values(), key=lambda x: x["subtotal"], reverse=True)
        for g in group_summary:
            g["pct"] = round(g["subtotal"] / total_nonzero * 100, 1)

        # Agrupar segmentos < 3% en "Otros" para el donut (la card muestra todos)
        DONUT_MIN_PCT = 3
        main_groups = [g for g in group_summary if g["pct"] >= DONUT_MIN_PCT]
        small_groups = [g for g in group_summary if g["pct"] < DONUT_MIN_PCT]
        donut_groups = main_groups[:]
        if small_groups:
            otros_subtotal = sum(g["subtotal"] for g in small_groups)
            otros_pct = round(otros_subtotal / total_nonzero * 100, 1)
            donut_groups.append(
                {
                    "name": "Otros",
                    "color": "#adb5bd",
                    "pk": None,
                    "subtotal": otros_subtotal,
                    "pct": otros_pct,
                }
            )

        context["group_summary"] = group_summary
        context["category_summary"] = group_summary
        context["donut_labels"] = [g["name"] for g in donut_groups]
        context["donut_data"] = [float(g["subtotal"]) for g in donut_groups]
        context["donut_colors"] = [g["color"] for g in donut_groups]
        context["donut_pks"] = [g["pk"] for g in donut_groups]

        # Acumulado diario — solo cuando hay mes específico (explícito o default)
        has_filters = any(
            key in self.request.GET
            for key in [
                "q",
                "month",
                "year",
                "category",
                "subcategory",
                "date_from",
                "date_to",
                "payment_method",
            ]
        )
        if has_filters:
            month_str = self.request.GET.get("month")
            year_str = self.request.GET.get("year")
        else:
            today = timezone.localdate()
            month_str = str(today.month)
            year_str = str(today.year)

        daily_labels = []
        daily_data = []
        if month_str and year_str:
            try:
                month_int = int(month_str)
                year_int = int(year_str)
                if 1 <= month_int <= 12 and 1900 <= year_int <= 2100:
                    import calendar

                    daily_rows = {
                        row["date"].day: float(row["subtotal"])
                        for row in qs.values("date")
                        .annotate(subtotal=Sum("amount_ars"))
                        .order_by("date")
                    }
                    num_days = calendar.monthrange(year_int, month_int)[1]
                    acum = 0
                    daily_bar_data = []
                    for d in range(1, num_days + 1):
                        day_amount = daily_rows.get(d, 0)
                        acum += day_amount
                        daily_labels.append(d)
                        daily_data.append(round(acum, 2))
                        daily_bar_data.append(round(day_amount, 2))
                    context["show_daily_chart"] = True
                    context["daily_bar_data"] = daily_bar_data
            except ValueError:
                pass

        context["daily_labels"] = daily_labels
        context["daily_data"] = daily_data
        context.setdefault("show_daily_chart", False)
        context.setdefault("daily_bar_data", [])

        # Barras apiladas mensuales — solo cuando hay año sin mes específico
        context["show_monthly_chart"] = False
        show_monthly = bool(year_str) and not bool(month_str)

        if show_monthly and year_str:
            try:
                year_int = int(year_str)
                if 1900 <= year_int <= 2100:
                    from django.db.models.functions import ExtractMonth

                    # Query: todos los gastos del año del usuario (sin otros filtros activos)
                    qs_year = Expense.objects.filter(
                        user=self.request.user, date__year=year_int
                    ).select_related("category", "category__parent")

                    # Totales por grupo para el año — elegir top N grupos
                    MAX_GROUPS = 6
                    group_year_totals = {}
                    for row in qs_year.values(
                        "category__parent_id",
                        "category__parent__name",
                        "category__parent__color",
                        "category_id",
                        "category__name",
                        "category__color",
                    ).annotate(subtotal=Sum("amount_ars")):
                        if row["category__parent_id"]:
                            gid = row["category__parent_id"]
                            gname = row["category__parent__name"]
                            gcolor = row["category__parent__color"] or "#6c757d"
                        else:
                            gid = row["category_id"]
                            gname = row["category__name"]
                            gcolor = row["category__color"] or "#6c757d"
                        if gid not in group_year_totals:
                            group_year_totals[gid] = {"name": gname, "color": gcolor, "subtotal": 0}
                        group_year_totals[gid]["subtotal"] += row["subtotal"]

                    sorted_groups = sorted(
                        group_year_totals.items(), key=lambda x: x[1]["subtotal"], reverse=True
                    )
                    top_groups = dict(sorted_groups[:MAX_GROUPS])
                    has_others = len(sorted_groups) > MAX_GROUPS

                    # Query: totales por mes y grupo
                    monthly_rows = (
                        qs_year.annotate(month=ExtractMonth("date"))
                        .values(
                            "month",
                            "category__parent_id",
                            "category__parent__color",
                            "category_id",
                            "category__color",
                        )
                        .annotate(subtotal=Sum("amount_ars"))
                    )

                    # Construir matriz [grupo][mes]
                    month_data = {gid: [0] * 12 for gid in top_groups}
                    others_by_month = [0] * 12
                    for row in monthly_rows:
                        m = row["month"] - 1  # 0-indexed
                        if row["category__parent_id"]:
                            gid = row["category__parent_id"]
                        else:
                            gid = row["category_id"]
                        if gid in top_groups:
                            month_data[gid][m] += float(row["subtotal"])
                        elif has_others:
                            others_by_month[m] += float(row["subtotal"])

                    month_names = [
                        "Ene",
                        "Feb",
                        "Mar",
                        "Abr",
                        "May",
                        "Jun",
                        "Jul",
                        "Ago",
                        "Sep",
                        "Oct",
                        "Nov",
                        "Dic",
                    ]
                    datasets = [
                        {
                            "label": top_groups[gid]["name"],
                            "data": month_data[gid],
                            "color": top_groups[gid]["color"],
                        }
                        for gid in top_groups
                    ]
                    if has_others:
                        datasets.append(
                            {"label": "Otros", "data": others_by_month, "color": "#adb5bd"}
                        )

                    context["monthly_labels"] = month_names
                    context["monthly_datasets"] = datasets
                    context["show_monthly_chart"] = True
            except ValueError:
                pass

        return context


class ExpenseCreateView(UserOwnedCreateView):
    model = Expense
    form_class = ExpenseForm
    template_name = "expenses/expense_form.html"
    success_url = reverse_lazy("expenses:list")

    def get_success_url(self):
        if self.request.GET.get("recurring") or self.request.POST.get("recurring"):
            return reverse_lazy("recurring:list")
        return super().get_success_url()

    def form_valid(self, form):
        response = super().form_valid(form)
        recurring_pk = self.request.GET.get("recurring") or self.request.POST.get("recurring")
        if recurring_pk:
            try:
                from apps.recurring.models import RecurringExpense

                rec = RecurringExpense.objects.get(pk=recurring_pk, user=self.request.user)
                rec.auto_deactivate_if_complete()
            except (RecurringExpense.DoesNotExist, ValueError):
                pass
        return response

    def get_success_message(self):
        obj = self.object
        return f"Gasto registrado: {obj.description} - {obj.formatted_amount}"

    def form_invalid(self, form):
        messages.error(
            self.request,
            "No pudimos guardar el gasto. Revisá los campos marcados. Monto, categoría y fecha son obligatorios.",
        )
        return super().form_invalid(form)

    def get_initial(self):
        initial = super().get_initial()
        recurring_pk = self.request.GET.get("recurring")
        if recurring_pk:
            try:
                from apps.recurring.models import RecurringExpense

                recurring = RecurringExpense.objects.get(pk=recurring_pk, user=self.request.user)
                initial["recurring"] = recurring.pk
                initial["category"] = recurring.category
                initial["description"] = recurring.name
            except (RecurringExpense.DoesNotExist, ValueError):
                pass

        duplicate_pk = self.request.GET.get("duplicate")
        if duplicate_pk:
            try:
                source = Expense.objects.get(pk=duplicate_pk, user=self.request.user)
                initial["category"] = source.category
                initial["description"] = source.description
                initial["amount"] = source.amount
                initial["currency"] = source.currency
                initial["exchange_rate"] = source.exchange_rate
                initial["payment_method"] = source.payment_method
            except (Expense.DoesNotExist, ValueError):
                pass

        # Precargar última cotización USD usada por el usuario
        if not initial.get("exchange_rate"):
            last_usd = (
                Expense.objects.filter(user=self.request.user, currency="USD")
                .order_by("-date", "-created_at")
                .values_list("exchange_rate", flat=True)
                .first()
            )
            if last_usd:
                initial["exchange_rate"] = last_usd

        return initial

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["categories_by_group"] = Category.get_categories_by_group(
            self.request.user, "EXPENSE"
        )
        recurring_pk = self.request.GET.get("recurring")
        if recurring_pk:
            try:
                from apps.recurring.models import RecurringExpense

                context["linked_recurring"] = RecurringExpense.objects.get(
                    pk=recurring_pk, user=self.request.user
                )
            except (RecurringExpense.DoesNotExist, ValueError):
                pass
        return context


class ExpenseUpdateView(UserOwnedUpdateView):
    model = Expense
    form_class = ExpenseForm
    template_name = "expenses/expense_form.html"
    success_url = reverse_lazy("expenses:list")

    def get_success_message(self):
        obj = self.object
        return f"Gasto actualizado: {obj.description} - {obj.formatted_amount}"

    def form_invalid(self, form):
        messages.error(
            self.request,
            "No pudimos guardar el gasto. Revisá los campos marcados.",
        )
        return super().form_invalid(form)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["categories_by_group"] = Category.get_categories_by_group(
            self.request.user, "EXPENSE"
        )
        return context


class ExpenseDeleteView(UserOwnedDeleteView):
    model = Expense
    template_name = "expenses/expense_confirm_delete.html"
    success_url = reverse_lazy("expenses:list")

    def get_success_message(self, obj):
        return f"Gasto '{obj.description}' eliminado correctamente."


class ExpenseDetailView(UserOwnedDetailView):
    model = Expense
    template_name = "expenses/expense_detail.html"
    context_object_name = "expense"

    def get_queryset(self):
        return super().get_queryset().select_related("category", "saving")


class ExpenseExportView(ExpenseListView):
    """Exporta los gastos filtrados como XLSX, respetando los mismos filtros que la lista."""

    def get(self, request, *args, **kwargs):
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill

        expenses = self.get_queryset().select_related("category", "category__parent")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Gastos"

        headers = [
            "Fecha",
            "Grupo",
            "Subcategoría",
            "Descripción",
            "Monto",
            "Moneda",
            "Tipo de cambio",
            "Monto ARS",
            "Método de pago",
        ]
        ws.append(headers)

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="1F4E79")
        center = Alignment(horizontal="center")
        right = Alignment(horizontal="right")

        for _col, cell in enumerate(ws[1], start=1):
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center

        column_widths = [12, 22, 22, 36, 14, 10, 14, 14, 18]
        for i, width in enumerate(column_widths, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

        payment_method_labels = dict(PaymentMethod.choices)

        for expense in expenses:
            cat = expense.category
            grupo = cat.parent.name if cat.parent else cat.name
            subcategoria = cat.name if cat.parent else ""
            ws.append(
                [
                    expense.date.strftime("%d/%m/%Y"),
                    grupo,
                    subcategoria,
                    expense.description,
                    float(expense.amount),
                    expense.currency,
                    float(expense.exchange_rate) if expense.exchange_rate else "",
                    float(expense.amount_ars),
                    payment_method_labels.get(expense.payment_method, "")
                    if expense.payment_method
                    else "",
                ]
            )

        for row in ws.iter_rows(min_row=2):
            for col_idx, cell in enumerate(row, start=1):
                if col_idx in (5, 7, 8):
                    cell.alignment = right

        today = timezone.localdate()
        filename = f"gastos {today.strftime('%d.%m.%Y')}.xlsx"
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response


def _build_import_preview(rows, user, categories):
    """
    Convierte lista de ImportedRow en lista de dicts para el template de preview.
    - Agrupa impuestos ARS en una sola fila.
    - Sugiere categoría por historial del usuario (descripción exacta) con prioridad
      sobre la sugerencia por nombre de la fila.
    - La fila de impuestos agrupada se marca como is_tax=True para que el template
      la deje deschekeada por defecto.
    """
    from decimal import Decimal as _D

    from .importers import ImportedRow

    # Agrupar impuestos ARS
    tax_total = _D("0")
    tax_date = None
    non_tax_rows = []
    for r in rows:
        if r.suggested_category_name == "Impuestos" and r.currency == "ARS":
            tax_total += r.amount
            if tax_date is None:
                tax_date = r.raw_date
        else:
            non_tax_rows.append(r)

    if tax_total > 0:
        non_tax_rows.append(
            ImportedRow(
                raw_date=tax_date,
                description="Impuestos tarjeta",
                amount=tax_total,
                currency="ARS",
                suggested_category_name="Impuestos",
            )
        )

    # Mapa nombre → pk para sugerencias por nombre (impuestos)
    cat_by_name = {c.name.lower(): c.pk for c in categories}

    # Mapa descripción → pk basado en historial del usuario
    descriptions = [r.description.strip() for r in non_tax_rows]
    history_map: dict[str, int] = {}
    if descriptions:
        past = (
            Expense.objects.filter(user=user, description__in=descriptions)
            .exclude(category__isnull=True)
            .values("description", "category_id")
            .order_by("-date")
        )
        for row in past:
            desc = row["description"]
            if desc not in history_map:
                history_map[desc] = row["category_id"]

    # Set de (fecha, descripción, monto) ya existentes para detectar duplicados
    existing_keys = set()
    if descriptions:
        existing = Expense.objects.filter(user=user, description__in=descriptions).values(
            "date", "description", "amount"
        )
        for row in existing:
            existing_keys.add((row["date"], row["description"], row["amount"]))

    # Set de nombres de recurrentes activos: filas Fijo/Cuota que coincidan se
    # vinculan al recurrente existente en lugar de crear uno nuevo (ver Fase 2/3).
    from apps.recurring.models import RecurringExpense

    existing_recurring_names = set()
    if descriptions:
        existing_recurring_names = set(
            RecurringExpense.objects.filter(
                user=user, is_active=True, name__in=descriptions
            ).values_list("name", flat=True)
        )

    preview_rows = []
    for r in non_tax_rows:
        desc = r.description.strip().replace("\n", " ").replace("\r", "")
        is_tax = r.suggested_category_name == "Impuestos"

        # Prioridad: historial > sugerencia por nombre
        suggested_pk = history_map.get(desc) or cat_by_name.get(
            r.suggested_category_name.lower(), ""
        )

        is_duplicate = not is_tax and (r.date_parsed, desc, r.amount) in existing_keys

        installment_match = r.installment_match
        preview_rows.append(
            {
                "date": r.raw_date,
                "description": desc,
                "amount": str(r.amount),
                "currency": r.currency,
                "suggested_category_pk": suggested_pk,
                "is_tax": is_tax,
                "is_duplicate": is_duplicate,
                "suggested_type": "punctual" if is_tax else r.suggested_type,
                "installment_current": installment_match[0] if installment_match else "",
                "installment_total": installment_match[1] if installment_match else "",
                "matches_existing_recurring": desc in existing_recurring_names,
            }
        )

    return preview_rows


def _import_preview_hash(preview_rows):
    """Hash determinístico del contenido del resumen, para atar el progreso guardado
    en localStorage a un PDF específico (mismo resumen re-subido = mismo hash)."""
    raw = "|".join(f"{r['date']}:{r['description']}:{r['amount']}" for r in preview_rows)
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()[:16]


class ExpenseImportView(LoginRequiredMixin, View):
    """Paso 1: el usuario sube el PDF y ve el preview de transacciones."""

    template_name = "expenses/expense_import.html"

    def get(self, request):
        categories = Category.get_expense_categories(request.user)
        return render(request, self.template_name, {"categories": categories})

    def post(self, request):
        from .importers import parse_macro_visa

        pdf_file = request.FILES.get("pdf_file")
        if not pdf_file:
            messages.error(request, "Seleccioná un archivo PDF.")
            return redirect("expenses:import")

        if not pdf_file.name.lower().endswith(".pdf"):
            messages.error(request, "El archivo debe ser un PDF.")
            return redirect("expenses:import")

        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
            for chunk in pdf_file.chunks():
                tmp.write(chunk)
            tmp_path = tmp.name

        try:
            rows = parse_macro_visa(tmp_path)
        except Exception:
            logger.exception("Error al parsear el PDF")
            messages.error(
                request,
                "No se pudo leer el archivo. Verificá que sea un resumen Visa Macro válido.",
            )
            return redirect("expenses:import")

        if not rows:
            messages.warning(request, "No se encontraron transacciones en el archivo.")
            return redirect("expenses:import")

        categories = Category.get_expense_categories(request.user)
        expense_groups = Category.get_groups(request.user, CategoryType.EXPENSE)
        preview_rows = _build_import_preview(rows, request.user, categories)
        return render(
            request,
            self.template_name,
            {
                "categories": categories,
                "expense_groups": expense_groups,
                "preview_rows": preview_rows,
                "preview_json": json.dumps(preview_rows),
                "total_rows": len(preview_rows),
                "has_usd_rows": any(r["currency"] == "USD" for r in preview_rows),
                "preview_hash": _import_preview_hash(preview_rows),
            },
        )


class ExpenseImportDebugView(LoginRequiredMixin, View):
    """Vista de debug: parsea un PDF fijo del servidor y muestra el preview sin upload.
    Solo disponible con DEBUG=True."""

    template_name = "expenses/expense_import.html"

    def get(self, request):
        from django.conf import settings

        if not settings.DEBUG:
            from django.http import Http404

            raise Http404

        from pathlib import Path

        from .importers import parse_macro_visa

        pdf_path = Path(settings.BASE_DIR) / "resumen_6_2026.pdf"

        if not pdf_path.exists():
            messages.error(request, f"PDF de debug no encontrado en {pdf_path}")
            return redirect("expenses:import")

        try:
            rows = parse_macro_visa(pdf_path)
        except Exception:
            logger.exception("Error al parsear PDF de debug")
            messages.error(request, "No se pudo leer el PDF de debug.")
            return redirect("expenses:import")

        categories = Category.get_expense_categories(request.user)
        expense_groups = Category.get_groups(request.user, CategoryType.EXPENSE)
        preview_rows = _build_import_preview(rows, request.user, categories)
        return render(
            request,
            self.template_name,
            {
                "categories": categories,
                "expense_groups": expense_groups,
                "preview_rows": preview_rows,
                "preview_json": json.dumps(preview_rows),
                "total_rows": len(preview_rows),
                "has_usd_rows": any(r["currency"] == "USD" for r in preview_rows),
                "preview_hash": _import_preview_hash(preview_rows),
            },
        )


class ExpenseImportConfirmView(LoginRequiredMixin, View):
    """Paso 2: el usuario confirma las filas y se crean los gastos."""

    def post(self, request):
        try:
            rows_data = json.loads(request.POST.get("rows_json", "[]"))
        except (ValueError, TypeError):
            messages.error(request, "Datos inválidos.")
            return redirect("expenses:import")

        from apps.recurring.models import RecurringExpense

        user = request.user
        categories = {str(c.pk): c for c in Category.get_expense_categories(user)}
        recurring_by_name = {
            r.name: r for r in RecurringExpense.objects.filter(user=user, is_active=True)
        }

        created = 0
        error_details: list[str] = []

        with transaction.atomic():
            for row in rows_data:
                if not row.get("include"):
                    continue

                row_desc = (row.get("description") or "(sin descripción)")[:60]

                cat_pk = str(row.get("category_pk", ""))
                category = categories.get(cat_pk)
                if not category:
                    error_details.append(f"{row_desc}: sin categoría válida")
                    continue

                raw_date = row.get("date", "")
                try:
                    d, m, y = raw_date.split(".")
                    expense_date = timezone.datetime(2000 + int(y), int(m), int(d)).date()
                except (ValueError, AttributeError):
                    error_details.append(f"{row_desc}: fecha inválida")
                    continue

                try:
                    amount = Decimal(row.get("amount", "0"))
                except InvalidOperation:
                    error_details.append(f"{row_desc}: monto inválido")
                    continue

                currency = row.get("currency", "ARS")
                description = row.get("description", "")[:255]

                if currency == "USD":
                    try:
                        exchange_rate = Decimal(str(row.get("exchange_rate") or "0"))
                    except InvalidOperation:
                        exchange_rate = Decimal("0")
                    if exchange_rate <= 0:
                        error_details.append(f"{row_desc}: falta cotización del dólar")
                        continue
                else:
                    exchange_rate = Decimal("1.0000")

                row_type = row.get("type", "punctual")
                recurring = None

                if row_type in ("fixed", "installment"):
                    recurring = recurring_by_name.get(description)
                    if recurring is None:
                        installment_total = None
                        starting_installment = None
                        if row_type == "installment":
                            try:
                                installment_total = int(row.get("installment_total") or 0)
                                starting_installment = int(row.get("installment_current") or 0)
                            except (TypeError, ValueError):
                                installment_total = None
                            if not installment_total or not starting_installment:
                                error_details.append(f"{row_desc}: datos de cuota inválidos")
                                continue

                        recurring = RecurringExpense.objects.create(
                            user=user,
                            name=description,
                            category=category,
                            due_day=min(expense_date.day, 28),
                            total_installments=installment_total,
                            starting_installment=starting_installment,
                            start_date=expense_date if row_type == "installment" else None,
                        )
                        recurring_by_name[description] = recurring

                Expense.objects.create(
                    user=user,
                    date=expense_date,
                    category=category,
                    description=description,
                    amount=amount,
                    currency=currency,
                    exchange_rate=exchange_rate,
                    recurring=recurring,
                )
                if recurring is not None:
                    recurring.auto_deactivate_if_complete()
                created += 1

        if created:
            messages.success(
                request,
                f"Se importaron {created} gasto{'s' if created != 1 else ''} correctamente.",
            )
        if error_details:
            n = len(error_details)
            detail_text = "; ".join(error_details[:5])
            if n > 5:
                detail_text += f"; y {n - 5} más"
            messages.warning(
                request,
                f"{n} fila{'s' if n != 1 else ''} no se pudieron importar: {detail_text}",
            )

        return redirect("expenses:list")
